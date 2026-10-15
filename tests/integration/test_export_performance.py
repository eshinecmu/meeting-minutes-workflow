"""PERF-04：固定约 2 万字输入，真实 DOCX 导出各次 ≤10s（§03 §8.1）。

冷：新 Python 子进程启动前计时，含导入、服务/字体报告初始化、读稿、
渲染、ZIP 序列化、落盘、manifest 与审计；不含预置数据库及输入文件。
热：同进程同服务第二版本完整导出，含同样的 I/O，不能命中幂等复用。
不清 OS 文件缓存，因此“冷”指进程冷启动，非整机/磁盘冷启动。
字体清单为确定性注入；无 ASR/LLM、LibreOffice、网络或模拟计时。
环境和精确规模写入 pytest 属性及 -s 输出；开发机通过不替代目标机验收。
§03 未规定导出超时错误码，此测试不套用 E3003。
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import platform
import subprocess
import sys
import time
from importlib.metadata import version
from pathlib import Path


FONTS = {"方正小标宋简体", "黑体", "楷体_GB2312", "仿宋_GB2312", "宋体", "Times New Roman"}
MEETING_ID = "export-perf-04"


def fixed_markdown() -> str:
    # 100 段，每段 200 个汉字；10 张表，每表 4 列、10 数据行加表头。
    paragraph = "会议要求各部门按时落实任务并完成检查" * 10
    assert len(paragraph) == 180
    paragraph += "各部门应及时报告工作进展并先明确责任分工"
    assert len(paragraph) == 200
    lines = ["# 年度工作会议纪要"]
    for section in range(1, 11):
        lines += [f"## 工作安排{section}"] + [paragraph] * 10
        lines += [f"表 {section} 任务分工", "| 任务 | 责任人 | 时限 | 状态 |",
                  "| --- | --- | --- | --- |"]
        lines += [f"| 检查{row} | 张三 | 本月底 | 待办 |" for row in range(1, 11)]
    return "\n\n".join(lines).replace("|\n\n|", "|\n|")


def _worker(data_dir: Path, cold_start: float) -> None:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
    from mmwf.app.services.export import ExportService
    from mmwf.infra.filestore import make_filestore

    store = make_filestore(MEETING_ID, data_dir, encrypt_enabled=False)
    service = ExportService(store, db_path=data_dir / "mmwf.db")
    samples = []
    for number, name in ((1, "cold"), (2, "warm")):
        start = cold_start if number == 1 else time.perf_counter()
        result = service.export_docx(number, installed_fonts=FONTS, convert_doc=False)
        elapsed = time.perf_counter() - start
        samples.append({"phase": name, "seconds": elapsed, "reused": result.reused,
                        "path": str(store.resolve(result.rel_path)), "sha256": result.sha256})
    environment = {"os": platform.platform(), "machine": platform.machine(),
                   "processor": platform.processor(), "logical_cpus": os.cpu_count(),
                   "python": sys.version, "executable": sys.executable,
                   "python_docx": version("python-docx"), "lxml": version("lxml"),
                   "font_probe": "injected installed-font set", "storage": str(data_dir)}
    print(json.dumps({"samples": samples, "environment": environment}, ensure_ascii=False))


def test_docx_export_within_ten_seconds_cold_and_warm(tmp_path, record_property):
    from docx import Document

    from mmwf.infra.audit import query
    from mmwf.infra.filestore import make_filestore
    from mmwf.infra.migrations.runner import migrate

    markdown = fixed_markdown()
    migrate(tmp_path / "mmwf.db")
    store = make_filestore(MEETING_ID, tmp_path, encrypt_enabled=False)
    for number in (1, 2):
        store.write_text(f"record/v{number}.md", markdown, stage="record")
    # 子进程第一条导出包含导入开销；计时器不替换。30s 仅作挂死保护。
    cold_start = time.perf_counter()
    completed = subprocess.run(
        [sys.executable, str(Path(__file__).resolve()), "--worker", str(tmp_path), str(cold_start)],
        capture_output=True, text=True, timeout=30, check=True,
    )
    report = json.loads(completed.stdout)
    report["input"] = {"characters": len(markdown), "utf8_bytes": len(markdown.encode("utf-8")),
                       "body_paragraphs": 100, "body_hanzi": 20000, "headings": 10,
                       "tables": 10, "table_data_rows": 100, "table_columns": 4,
                       "encryption": False, "audit": True}
    record_property("PERF-04", json.dumps(report, ensure_ascii=False))
    print("PERF-04 " + json.dumps(report, ensure_ascii=False))
    assert [s["phase"] for s in report["samples"]] == ["cold", "warm"]
    paths = set()
    for sample in report["samples"]:
        assert 0 < sample["seconds"] <= 10, sample
        assert sample["reused"] is False
        paths.add(sample["path"])
        payload = Path(sample["path"]).read_bytes()
        assert hashlib.sha256(payload).hexdigest() == sample["sha256"]
        document = Document(io.BytesIO(payload))
        assert len(document.tables) == 10
        assert all(len(t.rows) == 11 and len(t.columns) == 4 for t in document.tables)
        assert sum(len(p.text) == 200 for p in document.paragraphs) == 100
        assert "年度工作会议纪要" in document.paragraphs[0].text
        assert document.tables[-1].cell(10, 0).text == "检查 10"
    assert len(paths) == 2
    exports = [entry for entry in store.manifest() if entry.stage == "export"]
    assert len(exports) == 2
    records = query(tmp_path / "mmwf.db")
    assert len(records) == 2 and all(r.action == "export" and r.result == "success" for r in records)


if __name__ == "__main__":
    assert sys.argv[1] == "--worker"
    _worker(Path(sys.argv[2]), float(sys.argv[3]))
